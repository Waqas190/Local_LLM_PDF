import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import requests
import PyPDF2
import docx
import os

# Ollama API URL
url = "http://localhost:11434/v1/chat/completions"

# Function to extract text from a PDF
def extract_pdf_text(pdf_file_path):
    try:
        with open(pdf_file_path, 'rb') as file:
            reader = PyPDF2.PdfReader(file)
            text = ""
            for page in reader.pages:
                text += page.extract_text()
            return text
    except Exception as e:
        return str(e)

# Function to extract text from a PowerPoint (PPTX) file
def extract_pptx_text(pptx_file_path):
    try:
        presentation = Presentation(pptx_file_path)
        text = ""
        for slide in presentation.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text += shape.text
        return text
    except Exception as e:
        return str(e)

# Function to extract text from a Word (DOCX) file
def extract_docx_text(docx_file_path):
    try:
        doc = docx.Document(docx_file_path)
        text = ""
        for para in doc.paragraphs:
            text += para.text + '\n'
        return text
    except Exception as e:
        return str(e)

# Function to extract text from a TXT file
def extract_txt_text(txt_file_path):
    try:
        with open(txt_file_path, 'r', encoding='utf-8') as file:
            return file.read()
    except Exception as e:
        return str(e)

# Function to interact with Ollama API
def interact_with_ollama(text, question):
    try:
        messages = [
            {"role": "user", "content": f"Here is the text from the file:\n{text}\n\nPlease answer the following question: {question}"}
        ]
        response = requests.post(url, json={"model": "llama2", "messages": messages})
        
        if response.status_code == 200:
            return response.json()['choices'][0]['message']['content']
        else:
            return f"Error: API request failed with status code {response.status_code}"
    except Exception as e:
        return f"Error: {str(e)}"

# Function to handle file selection
def select_file():
    file_path = filedialog.askopenfilename(filetypes=[
        ("Supported Files", "*.pdf;*.pptx;*.docx;*.txt"),
        ("PDF Files", "*.pdf"),
        ("PowerPoint Files", "*.pptx"),
        ("Word Files", "*.docx"),
        ("Text Files", "*.txt")])
    
    if file_path:
        file_extension = file_path.split('.')[-1].lower()
        if file_extension == 'pdf':
            text = extract_pdf_text(file_path)
        elif file_extension == 'pptx':
            text = extract_pptx_text(file_path)
        elif file_extension == 'docx':
            text = extract_docx_text(file_path)
        elif file_extension == 'txt':
            text = extract_txt_text(file_path)
        else:
            messagebox.showerror("Error", "Unsupported file type.")
            return
        
        if text:
            question_entry.config(state=tk.NORMAL)
            ask_button.config(state=tk.NORMAL)
            global file_content
            file_content.set(text)
        else:
            messagebox.showerror("Error", "Failed to extract text from the file.")

# Function to handle asking questions
def ask_question():
    question = question_entry.get()
    if not question:
        messagebox.showwarning("Warning", "Please enter a question.")
        return
    
    text = file_content.get()
    if not text:
        messagebox.showwarning("Warning", "No file text available to query.")
        return
    
    answer = interact_with_ollama(text, question)
    result_text.delete(1.0, tk.END)
    result_text.insert(tk.END, answer)

# Function to clear results
def clear_results():
    question_entry.delete(0, tk.END)
    result_text.delete(1.0, tk.END)

# Main UI setup
root = tk.Tk()
root.title("File Chat with Ollama")
root.geometry("800x600")

style = ttk.Style()
style.theme_use("clam")
style.configure("TButton", font=("Helvetica", 12))
style.configure("TLabel", font=("Helvetica", 12))

select_button = ttk.Button(root, text="Select File", command=select_file)
select_button.pack(pady=10)

question_label = ttk.Label(root, text="Ask a question about the file:")
question_label.pack(pady=5)

question_entry = ttk.Entry(root, width=70, state=tk.DISABLED)
question_entry.pack(pady=5)

ask_button = ttk.Button(root, text="Ask", command=ask_question, state=tk.DISABLED)
ask_button.pack(pady=10)

result_label = ttk.Label(root, text="Answer:")
result_label.pack(pady=5)

result_text = tk.Text(root, width=90, height=15, wrap=tk.WORD)
result_text.pack(pady=10)

clear_button = ttk.Button(root, text="Clear", command=clear_results)
clear_button.pack(pady=5)

file_content = tk.StringVar()

root.mainloop()
